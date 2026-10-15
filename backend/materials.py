"""材料解析入口：原生结构优先，图片和无文字 PDF 再启用本地 OCR。"""

from __future__ import annotations

import importlib.metadata
import importlib.util
import os
import re
import threading
import zipfile
from dataclasses import dataclass
from email import policy
from email.parser import BytesParser
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from pypdf import PdfReader


IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}


def _version(package: str) -> str:
    try:
        return importlib.metadata.version(package)
    except importlib.metadata.PackageNotFoundError:
        return "not-installed"


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    status: str
    issue: str | None
    parser_name: str
    parser_version: str


class LocalOcrReader:
    """PaddleOCR 3.x/2.x 兼容适配器；模型仅在第一次 OCR 时加载。"""

    def __init__(self, cache_root: Path):
        self.cache_root = cache_root
        self._engine = None
        self._lock = threading.Lock()

    @property
    def available(self) -> bool:
        return (
            importlib.util.find_spec("paddleocr") is not None
            and importlib.util.find_spec("paddle") is not None
        )

    def _load(self):
        with self._lock:
            if self._engine is not None:
                return self._engine
            if not self.available:
                raise RuntimeError("本机尚未安装 PaddleOCR")
            self.cache_root.mkdir(parents=True, exist_ok=True)
            os.environ.setdefault("PADDLE_PDX_CACHE_HOME", str(self.cache_root / "paddlex"))
            os.environ.setdefault("PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK", "True")
            from paddleocr import PaddleOCR

            modern = {
                "text_detection_model_name": "PP-OCRv6_medium_det",
                "text_recognition_model_name": "PP-OCRv6_medium_rec",
                "use_doc_orientation_classify": False,
                "use_doc_unwarping": False,
                "use_textline_orientation": False,
                "enable_mkldnn": False,
                "engine_config": {
                    "paddle_static": {
                        "run_mode": "paddle",
                        "cpu_threads": 8,
                        "enable_new_ir": False,
                    }
                },
                "device": "cpu",
            }
            try:
                self._engine = PaddleOCR(**modern)
            except (TypeError, ValueError):
                self._engine = PaddleOCR(lang="ch", use_angle_cls=True, show_log=False)
            return self._engine

    @staticmethod
    def _plain(value: Any) -> Any:
        if isinstance(value, (dict, list, tuple, str, int, float, bool)) or value is None:
            return value
        candidate = getattr(value, "json", None)
        if candidate is not None:
            candidate = candidate() if callable(candidate) else candidate
            if isinstance(candidate, (dict, list, tuple)):
                return candidate
        candidate = getattr(value, "to_dict", None)
        return candidate() if callable(candidate) else value

    @classmethod
    def _modern_lines(cls, value: Any) -> list[tuple[str, float | None]]:
        value = cls._plain(value)
        if isinstance(value, dict) and isinstance(value.get("res"), dict):
            value = value["res"]
        if not isinstance(value, dict):
            return []
        texts = value.get("rec_texts") or value.get("texts")
        if not isinstance(texts, (list, tuple)):
            return []
        scores = value.get("rec_scores") or value.get("scores") or []
        return [
            (
                str(text),
                float(scores[index]) if index < len(scores) and scores[index] is not None else None,
            )
            for index, text in enumerate(texts)
        ]

    @staticmethod
    def _legacy_lines(value: Any) -> list[tuple[str, float | None]]:
        output: list[tuple[str, float | None]] = []

        def visit(node: Any) -> None:
            if (
                isinstance(node, (list, tuple))
                and len(node) == 2
                and isinstance(node[1], (list, tuple))
                and node[1]
                and isinstance(node[1][0], str)
            ):
                score = (
                    float(node[1][1])
                    if len(node[1]) > 1 and isinstance(node[1][1], (int, float))
                    else None
                )
                output.append((node[1][0], score))
                return
            if isinstance(node, (list, tuple)):
                for child in node:
                    visit(child)

        visit(value)
        return output

    def read(self, path: Path) -> str:
        engine = self._load()
        with self._lock:
            raw_pages = (
                list(engine.predict(input=str(path)))
                if hasattr(engine, "predict")
                else engine.ocr(str(path), cls=True)
            )
        pages = raw_pages if isinstance(raw_pages, list) else [raw_pages]
        output: list[str] = []
        for page_index, page in enumerate(pages, start=1):
            lines = self._modern_lines(page) or self._legacy_lines(page)
            readable = [re.sub(r"\s+", " ", text).strip() for text, _ in lines if text.strip()]
            if readable:
                output.append(f"## 第{page_index}页（OCR）\n" + "\n".join(readable))
        if not output:
            raise RuntimeError("OCR 没有识别出文字")
        return "\n\n".join(output)


_OCR = LocalOcrReader(
    Path(os.getenv("ZHOUXU_MODEL_CACHE", r"C:\Users\Public\ZhouXu\models"))
)


def extract_material(filename: str, payload: bytes, stored_path: Path) -> ExtractionResult:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in {".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm"}:
            for encoding in ("utf-8-sig", "gb18030"):
                try:
                    return ExtractionResult(
                        payload.decode(encoding), "read", None, "native-text", "1"
                    )
                except UnicodeDecodeError:
                    pass
            return ExtractionResult(
                payload.decode("utf-8", errors="replace"),
                "read_with_replacement",
                "部分字符使用替代符读取",
                "native-text",
                "1",
            )

        if suffix == ".eml":
            message = BytesParser(policy=policy.default).parsebytes(payload)
            lines = [
                f"From: {message.get('From', '')}",
                f"To: {message.get('To', '')}",
                f"Cc: {message.get('Cc', '')}",
                f"Date: {message.get('Date', '')}",
                f"Subject: {message.get('Subject', '')}",
            ]
            body = message.get_body(preferencelist=("plain",)) if message.is_multipart() else message
            if body:
                try:
                    lines.append(body.get_content())
                except Exception:
                    lines.append(str(body.get_payload(decode=True) or b"", "utf-8", "replace"))
            return ExtractionResult("\n".join(lines), "read", None, "python-email", "stdlib")

        if suffix == ".xlsx":
            workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in workbook.worksheets:
                lines.append(f"## 工作表：{sheet.title}")
                for row_no, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    values = [str(value).strip() for value in row if value not in (None, "")]
                    if values:
                        lines.append(f"第{row_no}行：" + " | ".join(values))
            workbook.close()
            return ExtractionResult(
                "\n".join(lines), "read", None, "openpyxl", _version("openpyxl")
            )

        if suffix == ".pdf":
            reader = PdfReader(BytesIO(payload))
            pages = [
                f"## 第{index}页\n{page.extract_text() or ''}"
                for index, page in enumerate(reader.pages, start=1)
            ]
            text = "\n\n".join(pages).strip()
            if text:
                return ExtractionResult(text, "read", None, "pypdf", _version("pypdf"))
            text = _OCR.read(stored_path)
            return ExtractionResult(
                text,
                "read_ocr",
                "文字由本地 OCR 生成，关键内容需要对照原件",
                "paddleocr",
                _version("paddleocr"),
            )

        if suffix == ".docx":
            with zipfile.ZipFile(BytesIO(payload)) as archive:
                xml = archive.read("word/document.xml")
            root = ElementTree.fromstring(xml)
            namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            paragraphs = [
                "".join(node.text or "" for node in paragraph.iter(namespace + "t"))
                for paragraph in root.iter(namespace + "p")
            ]
            return ExtractionResult(
                "\n".join(value for value in paragraphs if value.strip()),
                "read",
                None,
                "docx-xml",
                "1",
            )

        if suffix in IMAGE_SUFFIXES:
            text = _OCR.read(stored_path)
            return ExtractionResult(
                text,
                "read_ocr",
                "文字由本地 OCR 生成，关键内容需要对照原件",
                "paddleocr",
                _version("paddleocr"),
            )

        return ExtractionResult(
            "", "unread", f"暂不支持读取 {suffix or '无扩展名'} 文件，原件已经保存", "none", "1"
        )
    except Exception as exc:
        return ExtractionResult(
            "", "unread", f"文件读取失败：{type(exc).__name__}：{exc}", "failed", "1"
        )
